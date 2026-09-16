# Shared imports work in both source (core/) and installed flat layouts.
import sys as _sys
from pathlib import Path as _Path
_root = _Path(__file__).resolve().parents[2]
_sys.path.insert(0, str(_root / "core" if (_root / "core").is_dir() else _root))
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import warnings
import zipfile

from docx import Document

from vexi_foundation.contracts import Scope
from vexi_foundation.workspace import DocumentRecord
from vexi_foundation.document_lifecycle import LocalDocumentLifecycleStore, LifecycleError
from vexi_foundation.docx_provider import (
    ApprovalRequirement, DocxError, DocxLimits, FieldKind,
    HeadlessDocxProvider, HybridDocxEditController,
    ReplaceTextCommand, StoryScope,
)


def build_docx(split=False, duplicate=False, include_header=True):
    doc = Document()
    if split:
        p = doc.add_paragraph()
        p.add_run("Client: AC")
        p.add_run("ME")
    else:
        doc.add_paragraph("Client: ACME")
    if duplicate:
        doc.add_paragraph("Second ACME")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Label"
    table.cell(0, 1).text = "Value"
    if include_header:
        doc.sections[0].header.paragraphs[0].text = "Header ACME"
    buf = BytesIO(); doc.save(buf); return buf.getvalue()


def rewrite_zip(source, transform=None, add=None, duplicate=None):
    inp = zipfile.ZipFile(BytesIO(source), "r")
    out = BytesIO()
    with inp, zipfile.ZipFile(out, "w") as z:
        for info in inp.infolist():
            data = inp.read(info)
            if transform and info.filename in transform:
                data = transform[info.filename](data)
            z.writestr(info, data)
        for name, data in (add or {}).items():
            z.writestr(name, data)
        if duplicate:
            name, data = duplicate
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                z.writestr(name, data)
    return out.getvalue()


class ProviderTests(unittest.TestCase):
    def test_inspect_real_docx(self):
        p = HeadlessDocxProvider(); info = p.inspect(build_docx())
        self.assertGreaterEqual(info.paragraph_count, 3)
        self.assertGreaterEqual(info.table_count, 1)
        self.assertIn("word/document.xml", info.story_parts)

    def test_exact_replacement_across_word_runs(self):
        source = build_docx(split=True)
        provider = HeadlessDocxProvider()
        cmd = ReplaceTextCommand("ACME", "Umbra", 1, FieldKind.NON_MATERIAL_TEXT)
        plan = provider.plan(source, cmd, trusted_schema=True)
        self.assertEqual(plan.approval, ApprovalRequirement.AUTO_WITH_RECOVERY)
        output = provider.apply(source, cmd, plan, trusted_schema=True)
        reopened = Document(BytesIO(output))
        self.assertIn("Client: Umbra", "\n".join(p.text for p in reopened.paragraphs))

    def test_occurrence_mismatch_fails_closed(self):
        provider = HeadlessDocxProvider(); source = build_docx(duplicate=True)
        with self.assertRaisesRegex(DocxError, "occurrence_mismatch"):
            provider.plan(source, ReplaceTextCommand("ACME", "X", 1), trusted_schema=False)

    def test_table_cell_text_is_supported(self):
        provider = HeadlessDocxProvider(); source = build_docx()
        cmd = ReplaceTextCommand("Value", "Updated", 1, FieldKind.NON_MATERIAL_TEXT)
        plan = provider.plan(source, cmd, trusted_schema=True)
        output = provider.apply(source, cmd, plan, trusted_schema=True)
        doc = Document(BytesIO(output))
        self.assertEqual(doc.tables[0].cell(0, 1).text, "Updated")

    def test_header_scope(self):
        provider = HeadlessDocxProvider(); source = build_docx()
        cmd = ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT, StoryScope.HEADERS)
        plan = provider.plan(source, cmd, trusted_schema=True)
        output = provider.apply(source, cmd, plan, trusted_schema=True)
        doc = Document(BytesIO(output))
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "Header NOVA")
        self.assertEqual(doc.paragraphs[0].text, "Client: ACME")

    def test_untrusted_low_risk_still_requires_confirmation(self):
        plan = HeadlessDocxProvider().plan(
            build_docx(), ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT),
            trusted_schema=False)
        self.assertEqual(plan.approval, ApprovalRequirement.OWNER_CONFIRMATION)

    def test_amount_requires_confirmation_even_trusted(self):
        plan = HeadlessDocxProvider().plan(
            build_docx(), ReplaceTextCommand("ACME", "1000 EUR", 1, FieldKind.AMOUNT),
            trusted_schema=True)
        self.assertEqual(plan.approval, ApprovalRequirement.OWNER_CONFIRMATION)

    def test_macro_document_blocked(self):
        source = rewrite_zip(build_docx(), add={"word/vbaProject.bin": b"x"})
        with self.assertRaisesRegex(DocxError, "macro_enabled_unsupported"):
            HeadlessDocxProvider().inspect(source)

    def test_signed_document_blocked(self):
        source = rewrite_zip(build_docx(), add={"_xmlsignatures/sig1.xml": b"<x/>"})
        with self.assertRaisesRegex(DocxError, "signed_document_unsupported"):
            HeadlessDocxProvider().inspect(source)

    def test_unsafe_member_path_blocked(self):
        source = rewrite_zip(build_docx(), add={"../escape.bin": b"x"})
        with self.assertRaisesRegex(DocxError, "unsafe_member_path"):
            HeadlessDocxProvider().inspect(source)

    def test_duplicate_member_blocked(self):
        source = build_docx()
        source = rewrite_zip(source, duplicate=("word/document.xml", b"<x/>"))
        with self.assertRaisesRegex(DocxError, "duplicate_member"):
            HeadlessDocxProvider().inspect(source)

    def test_dtd_blocked(self):
        def inject(raw): return b'<!DOCTYPE x [<!ENTITY y "z">]>' + raw
        source = rewrite_zip(build_docx(), transform={"word/document.xml": inject})
        with self.assertRaisesRegex(DocxError, "dtd_entity_unsupported"):
            HeadlessDocxProvider().inspect(source)

    def test_tracked_changes_block_edit(self):
        def inject(raw):
            marker = b"<w:body>"
            return raw.replace(marker, marker + b'<w:ins w:id="1" w:author="x"/>', 1)
        source = rewrite_zip(build_docx(), transform={"word/document.xml": inject})
        with self.assertRaisesRegex(DocxError, "tracked_changes_unsupported"):
            HeadlessDocxProvider().plan(source, ReplaceTextCommand("ACME", "NOVA", 1))

    def test_member_limit(self):
        provider = HeadlessDocxProvider(DocxLimits(max_members=1))
        with self.assertRaisesRegex(DocxError, "too_many_members"):
            provider.inspect(build_docx())

    def test_non_target_member_preserved(self):
        source = rewrite_zip(build_docx(), add={"word/media/synthetic.bin": b"unchanged-media"})
        p = HeadlessDocxProvider(); cmd = ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT)
        out = p.apply(source, cmd, p.plan(source, cmd, trusted_schema=True), trusted_schema=True)
        with zipfile.ZipFile(BytesIO(out)) as z:
            self.assertEqual(z.read("word/media/synthetic.bin"), b"unchanged-media")


class ControllerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.scope = Scope("tenant-a", "project-a")
        self.store = LocalDocumentLifecycleStore(self.root / "state")
        self.record = DocumentRecord("docx-1", self.scope, "draft.docx", "DOCX", (), None, "", "")
        self.state = self.store.register(self.record, initial_content=build_docx(), now=1.0)

    def tearDown(self): self.tmp.cleanup()

    def test_low_risk_trusted_auto_with_recovery(self):
        c = HybridDocxEditController(self.store, schema_trust_resolver=lambda *a: True)
        result = c.execute("docx-1", self.scope,
            ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT, change_summary="label"),
            expected_generation=self.state.generation, now=2.0)
        self.assertEqual(result.code, "confirmed_success")
        self.assertEqual(len(result.state.versions), 1)  # pre-edit recovery version
        self.assertIsNone(result.post_version_id)
        self.assertIn("NOVA", Document(BytesIO(self.store.read_working("docx-1", self.scope))).paragraphs[0].text)

    def test_material_waits_without_mutation(self):
        before = self.store.read_working("docx-1", self.scope)
        c = HybridDocxEditController(self.store)
        result = c.execute("docx-1", self.scope,
            ReplaceTextCommand("ACME", "1000 EUR", 1, FieldKind.AMOUNT),
            expected_generation=self.state.generation, now=2.0)
        self.assertEqual(result.code, "awaiting_approval")
        self.assertEqual(self.store.load("docx-1", self.scope).generation, self.state.generation)
        self.assertEqual(self.store.read_working("docx-1", self.scope), before)

    def test_material_approval_is_bound_to_fingerprint(self):
        seen = {}
        def verifier(token, fp): seen["token"], seen["fp"] = token, fp; return token == "ok"
        c = HybridDocxEditController(self.store, approval_verifier=verifier, schema_trust_resolver=lambda *a: True)
        cmd = ReplaceTextCommand("ACME", "1000 EUR", 1, FieldKind.AMOUNT, change_summary="amount")
        result = c.execute("docx-1", self.scope, cmd, expected_generation=1,
                           approval_token="ok", now=2.0)
        self.assertEqual(result.code, "confirmed_success")
        self.assertEqual(seen["fp"], result.approval_fingerprint)
        self.assertEqual(len(result.state.versions), 2)  # recovery + approved material result
        self.assertIsNotNone(result.post_version_id)

    def test_bad_approval_does_not_create_recovery_or_mutate(self):
        c = HybridDocxEditController(self.store, approval_verifier=lambda t, f: False, schema_trust_resolver=lambda *a: True)
        with self.assertRaisesRegex(DocxError, "approval_invalid"):
            c.execute("docx-1", self.scope,
                ReplaceTextCommand("ACME", "1000", 1, FieldKind.AMOUNT),
                expected_generation=1, approval_token="bad", now=2.0)
        self.assertEqual(self.store.load("docx-1", self.scope).generation, 1)

    def test_stale_document_blocked(self):
        stale = self.store.mark_stale("docx-1", self.scope, expected_generation=1, reason_ref="external", now=2)
        with self.assertRaisesRegex(DocxError, "stale_document"):
            HybridDocxEditController(self.store, schema_trust_resolver=lambda *a: True).execute("docx-1", self.scope,
                ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT),
                expected_generation=stale.generation)

    def test_generation_conflict_before_plan(self):
        # Rebuilding DOCX can change ZIP timestamps across the 2-second boundary.
        original = self.store.read_working("docx-1", self.scope)
        self.store.autosave("docx-1", self.scope, content=original, expected_generation=1, now=2)
        # identical autosave is a no-op, so force change through stale marker
        state = self.store.mark_stale("docx-1", self.scope, expected_generation=1, reason_ref="x", now=3)
        with self.assertRaisesRegex(LifecycleError, "generation_conflict"):
            HybridDocxEditController(self.store, schema_trust_resolver=lambda *a: True).execute("docx-1", self.scope,
                ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT),
                expected_generation=1)
        self.assertGreater(state.generation, 1)

    def test_schema_trust_is_resolved_by_controller_not_command(self):
        c = HybridDocxEditController(self.store)
        result = c.execute("docx-1", self.scope,
            ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT),
            expected_generation=1, now=2)
        self.assertEqual(result.code, "awaiting_approval")
        self.assertEqual(self.store.load("docx-1", self.scope).generation, 1)

    def test_post_persist_verification_failure_rolls_back(self):
        class FailInspect(HeadlessDocxProvider):
            calls = 0
            def inspect(self, source):
                self.calls += 1
                if self.calls >= 1:
                    raise DocxError("synthetic_verify_failure")
                return super().inspect(source)
        original = self.store.read_working("docx-1", self.scope)
        c = HybridDocxEditController(self.store, provider=FailInspect(),
                                     schema_trust_resolver=lambda *a: True)
        with self.assertRaisesRegex(DocxError, "edit_failed_recovered"):
            c.execute("docx-1", self.scope,
                ReplaceTextCommand("ACME", "NOVA", 1, FieldKind.NON_MATERIAL_TEXT),
                expected_generation=1, now=2)
        self.assertEqual(self.store.read_working("docx-1", self.scope), original)

    def test_non_docx_record_blocked(self):
        other = DocumentRecord("note", self.scope, "a.md", "NOTE", (), None, "", "")
        self.store.register(other, initial_content=b"ACME", now=1)
        with self.assertRaisesRegex(DocxError, "document_type_not_docx"):
            HybridDocxEditController(self.store).execute("note", self.scope,
                ReplaceTextCommand("ACME", "NOVA", 1), expected_generation=1)


if __name__ == "__main__": unittest.main()
